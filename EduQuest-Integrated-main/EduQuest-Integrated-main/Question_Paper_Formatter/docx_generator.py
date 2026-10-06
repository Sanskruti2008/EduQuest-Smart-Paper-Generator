"""
docx_generator.py
-----------------
Generates the question-paper body DOCX.

Layout rules reverse-engineered from finaltemplete.docx and the institution
reference photo (Government Polytechnic Nashik exam paper body):

    Instructions:
    1. <instruction text>
    2. <instruction text>
    - - - - - - - - - - - - - - - - -

    Q. 1  Attempt any FIVE                                          (10)
    CO1  a)  Define operating system.                                (2)
    CO1  b)  Explain the concept of kernel in an operating system.   (2)

    Q. 2  Attempt any THREE                                         (12)
    CO1  a)  State the principles...                                 (4)

Layout rules from finaltemplete.docx:
  - Font:        Times New Roman, 10 pt body, 10 pt headings (bold for Q. headers)
  - Page:        A4 (8.27 × 11.69 in), margins L=1.25 R=1.0 T=0.75 B=0.75
  - Instructions: plain paragraphs, bold underlined heading, numbered list
  - Divider:     dashed line paragraph (matches template's "---" dividers)
  - Section hdr: "Q. 1  Attempt any FIVE" + right-tab "(10)", bold
  - Questions:   4-column bordered table  | CO | label | text | marks |
  - Col widths:  CO=0.38in  label=0.22in  text=fill  marks=0.40in
  - Cell margins: top/bottom=18twips; CO/label tight L/R; text wider
  - tblGrid + tblW set explicitly (prevents equal-width column reflow)
  - tblInd=0, tblCellSpacing=0 (tight, no inter-cell gaps)
  - Para spacing inside cells: sb=0 sa=1pt ls=1.0
  - Section space_before=8pt
  - NO institutional header generated
"""

import os
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches, Pt


# ---------------------------------------------------------------------------
# Layout constants — derived from finaltemplete.docx inspection
# ---------------------------------------------------------------------------

BODY_FONT    = "Times New Roman"   # institution body font
BODY_SIZE    = Pt(10)              # body text size
HEADING_SIZE = Pt(10)              # section header (bold)

# Page margins — match institution template
MARGIN_LEFT   = Inches(1.25)
MARGIN_RIGHT  = Inches(1.0)
MARGIN_TOP    = Inches(0.75)
MARGIN_BOTTOM = Inches(0.75)

# Column widths (inches) — CO narrow, label very narrow, text fills, marks narrow
COL_CO    = Inches(0.38)
COL_LABEL = Inches(0.22)
COL_MARKS = Inches(0.40)
# COL_TEXT  = printable_width - COL_CO - COL_LABEL - COL_MARKS (calculated per doc)

# Cell inner padding (twips — 1 in = 1440 twips)
CELL_TOP    = 18    # 0.013 in
CELL_BOTTOM = 18

# Per-column horizontal padding — metadata columns stay compact
CELL_MARGINS = {
    0: (36, 12),   # CO:      narrow left, minimal right gap to label
    1: (12, 12),   # label:   tight on both sides
    2: (36, 48),   # text:    normal left, wider right gap before marks
    3: (12, 36),   # marks:   tight left, normal right
}

# Paragraph spacing
INSTR_SB   = Pt(0)
INSTR_SA   = Pt(0)
INSTR_LS   = 1.0

ROW_SB     = Pt(0)
ROW_SA     = Pt(1)
ROW_LS     = 1.0

SEC_SB     = Pt(8)   # space before each Q. section
SEC_SA     = Pt(2)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get_subquestion_label(index: int) -> str:
    """
    Convert 0-based index to sub-question label.
    0 → 'a)', 1 → 'b)', …, 25 → 'z)', 26 → 'aa)', …
    """
    label = ""
    idx = index
    while True:
        label = chr(ord("a") + (idx % 26)) + label
        idx = idx // 26 - 1
        if idx < 0:
            break
    return f"{label})"


def _width_to_twips(width) -> int:
    """Convert a Length object or EMU int to twips for raw OOXML width attributes."""
    if hasattr(width, "twips"):
        return int(width.twips)
    return int(Emu(width).twips)


def _apply_cell_margins(cell, col_idx=0,
                        top=CELL_TOP, bottom=CELL_BOTTOM):
    left, right = CELL_MARGINS.get(col_idx, (36, 36))
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcMar")):
        tcPr.remove(old)
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top    w:w="{top}"    w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left   w:w="{left}"   w:type="dxa"/>'
        f'<w:right  w:w="{right}"  w:type="dxa"/>'
        f"</w:tcMar>"
    )
    tcPr.append(tcMar)


def _make_bordered(table):
    """Apply thin single-line borders matching the reference template grid."""
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top    w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left   w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right  w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        "</w:tblBorders>"
    )
    tblPr.append(tblBorders)


def _set_table_column_widths(table, col_widths):
    """
    Force explicit tblGrid + tblW + per-cell tcW widths.

    python-docx defaults tblGrid to equal-width columns when rows are added;
    without this, CO / label / marks columns expand to ~25 % of page width each.
    """
    widths_dxa = [_width_to_twips(w) for w in col_widths]
    total_dxa = sum(widths_dxa)

    tbl = table._tbl
    tblPr = tbl.tblPr

    # Full table width
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(parse_xml(
        f'<w:tblW {nsdecls("w")} w:w="{total_dxa}" w:type="dxa"/>'
    ))

    # Explicit column grid — must match cell count
    for old in tbl.findall(qn("w:tblGrid")):
        tbl.remove(old)
    grid_cols = "".join(
        f'<w:gridCol w:w="{w}"/>' for w in widths_dxa
    )
    tbl.insert(1, parse_xml(
        f'<w:tblGrid {nsdecls("w")}>{grid_cols}</w:tblGrid>'
    ))

    # Per-cell width enforcement
    for row in table.rows:
        for c_idx, w_dxa in enumerate(widths_dxa):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            tcPr = cell._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:tcW")):
                tcPr.remove(old)
            tcPr.append(parse_xml(
                f'<w:tcW {nsdecls("w")} w:w="{w_dxa}" w:type="dxa"/>'
            ))


def _set_fixed_layout(table):
    """
    Force Word to honour explicit column widths (tblLayout w:type="fixed")
    and remove default inter-cell spacing and indent.
    Without tblLayout=fixed, Word can auto-resize columns when content is long,
    breaking the alignment guarantee.
    """
    tblPr = table._tbl.tblPr

    # Fixed layout — columns never reflow
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    tblPr.append(parse_xml(
        f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'
    ))

    # Zero inter-cell spacing and table indent
    for tag in ("w:tblCellSpacing", "w:tblInd"):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    tblPr.append(parse_xml(
        f'<w:tblCellSpacing {nsdecls("w")} w:w="0" w:type="dxa"/>'
    ))
    tblPr.append(parse_xml(
        f'<w:tblInd {nsdecls("w")} w:w="0" w:type="dxa"/>'
    ))


# Kept for backward compatibility — _set_fixed_layout now handles both
_zero_table_spacing = _set_fixed_layout


def _set_para_fmt(para,
                  sb=ROW_SB, sa=ROW_SA, ls=ROW_LS,
                  align=WD_ALIGN_PARAGRAPH.LEFT):
    fmt = para.paragraph_format
    fmt.space_before = sb
    fmt.space_after  = sa
    fmt.line_spacing = ls
    para.alignment   = align


def _run(para, text, bold=False, font=BODY_FONT, size=BODY_SIZE):
    r = para.add_run(text)
    r.bold       = bold
    r.font.name  = font
    r.font.size  = size
    return r


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def generate_docx(formatted_paper: dict, output_path: str) -> None:
    """
    Generate the question-paper body DOCX.

    Args:
        formatted_paper: dict produced by paper_formatter.validate_and_format_paper()
        output_path:     destination file path (.docx)
    """
    doc = Document()

    # ── Page: A4 portrait, institution margins ───────────────────────────────
    page_sec = doc.sections[0]
    page_sec.page_width    = Inches(8.27)
    page_sec.page_height   = Inches(11.69)
    page_sec.left_margin   = MARGIN_LEFT
    page_sec.right_margin  = MARGIN_RIGHT
    page_sec.top_margin    = MARGIN_TOP
    page_sec.bottom_margin = MARGIN_BOTTOM

    # Printable width used for tab stops and table widths
    printable_w = page_sec.page_width - page_sec.left_margin - page_sec.right_margin

    col_co    = COL_CO
    col_label = COL_LABEL
    col_marks = COL_MARKS
    col_text  = printable_w - col_co - col_label - col_marks

    # ── Default paragraph style (Times New Roman 10pt) ───────────────────────
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE

    # ====================================================================== #
    # INSTRUCTIONS BLOCK                                                       #
    # ====================================================================== #
    instructions = formatted_paper.get("instructions", [])
    if instructions:
        # Heading: "Instructions:" — bold, underlined
        p_head = doc.add_paragraph()
        _set_para_fmt(p_head, sb=Pt(0), sa=Pt(1), ls=INSTR_LS)
        r_head = _run(p_head, "Instructions:", bold=True)
        r_head.underline = True

        # Numbered instruction items
        for idx, inst in enumerate(instructions, start=1):
            p_inst = doc.add_paragraph()
            _set_para_fmt(p_inst, sb=INSTR_SB, sa=INSTR_SA, ls=INSTR_LS)
            # Preserve faculty-supplied numbering, or prepend index
            text = inst if re.match(r"^\d+[\.\)]\s*", inst) else f"{idx}. {inst}"
            _run(p_inst, text)

        # Divider — dashed line matching the template's separator style
        p_div = doc.add_paragraph()
        _set_para_fmt(p_div, sb=Pt(2), sa=Pt(2), ls=INSTR_LS)
        r_div = _run(p_div, "- " * 40)   # matches template's "- - - - - -" style
        r_div.font.size = Pt(8)           # smaller so it reads as a divider

    # ====================================================================== #
    # QUESTION SECTIONS                                                        #
    # ====================================================================== #
    sections = formatted_paper.get("sections", {})

    for sec_name, sec_data in sections.items():
        label           = sec_data.get("section_label", sec_name)
        title           = (sec_data.get("title") or "").strip()
        attempted_marks = sec_data.get("attempted_marks")
        questions       = sec_data.get("questions", [])

        # ── Section header: "Q. 1  Attempt any FIVE\t(10)" ──────────────────
        p_sec = doc.add_paragraph()
        _set_para_fmt(p_sec, sb=SEC_SB, sa=SEC_SA, ls=1.0)
        # Right tab stop at printable width — pushes marks to right edge
        p_sec.paragraph_format.tab_stops.add_tab_stop(
            printable_w, WD_TAB_ALIGNMENT.RIGHT
        )

        header_body = f"{label}  {title}" if title else label
        _run(p_sec, header_body, bold=True)

        if attempted_marks is not None:
            _run(p_sec, f"\t({attempted_marks})", bold=True)

        if not questions:
            continue

        # ── 4-column bordered question table ────────────────────────────────
        # | CO (narrow) | label (narrow) | question text (fill) | marks (narrow) |
        tbl = doc.add_table(rows=0, cols=4, style="Table Grid")
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl.autofit   = False
        _make_bordered(tbl)
        _zero_table_spacing(tbl)
        _set_fixed_layout(tbl)

        col_widths = [col_co, col_label, col_text, col_marks]

        for q_idx, q in enumerate(questions):
            row = tbl.add_row()

            # Keep entire question row together — no mid-question page break
            trPr = row._tr.get_or_add_trPr()
            for old in trPr.findall(qn("w:cantSplit")):
                trPr.remove(old)
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            for c_idx, w in enumerate(col_widths):
                cell = row.cells[c_idx]
                cell.width = w
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                _apply_cell_margins(cell, col_idx=c_idx)

            co_text   = (q.get("course_outcome") or q.get("co") or "").strip()
            sub_label = get_subquestion_label(q_idx)
            q_text    = (q.get("question_text") or "").strip()
            q_marks   = q.get("marks")

            # C0 — CO (left-aligned)
            p0 = row.cells[0].paragraphs[0]
            _set_para_fmt(p0)
            _run(p0, co_text)

            # C1 — sub-label a) b) c)… (left-aligned)
            p1 = row.cells[1].paragraphs[0]
            _set_para_fmt(p1)
            _run(p1, sub_label)

            # C2 — question text (left-aligned, wraps naturally inside column)
            p2 = row.cells[2].paragraphs[0]
            _set_para_fmt(p2)
            _run(p2, q_text)

            # C3 — marks (right-aligned, fixed column)
            p3 = row.cells[3].paragraphs[0]
            _set_para_fmt(p3, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _run(p3, f"({q_marks})" if q_marks is not None else "")

        # Enforce non-equal column widths after all rows are populated
        _set_table_column_widths(tbl, col_widths)

    # ── Save ─────────────────────────────────────────────────────────────────
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir)
    doc.save(output_path)
